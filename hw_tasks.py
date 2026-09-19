# -*- coding: utf-8 -*-
"""作业解题器 · 任务管理：TaskState / TaskManager / 状态持久化 + 三个 worker。"""

import os
import json
import time
import copy
import threading
import uuid
from dataclasses import dataclass, field
from typing import Optional

import pymupdf as fitz
from openai import OpenAI

import hw_core
from hw_solve import (solve_page, solve_batch_office, apply_solution_to_page,
                      build_solution_markdown)
from hw_preview import (render_preview_only, _build_docx_preview_html,
                        _build_pptx_preview_html)

# Word / PPT 依赖为可选；未安装时这些名字为 None，
# 相关 worker 仅在 HAS_OFFICE 时才会被调用。
if hw_core.HAS_OFFICE:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    from pptx import Presentation
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except Exception:
        MSO_SHAPE_TYPE = None
else:
    Document = qn = OxmlElement = Paragraph = Presentation = None
    MSO_SHAPE_TYPE = None

_STATE_SAVE_TS = {}
_STATE_SAVE_LOCK = threading.RLock()


# ============================================================
# 任务管理
# ============================================================
@dataclass
class TaskState:
    task_id: str
    kind: str
    src_path: str
    src_name: str
    out_dir: str
    work_dir: str
    created_at: float = field(default_factory=time.time)
    status: str = "queued"
    current: int = 0
    total: int = 0
    label: str = ""
    log: list = field(default_factory=list)
    output_files: list = field(default_factory=list)
    preview_images: list = field(default_factory=list)
    preview_html: str = ""
    error: str = ""
    current_size: int = 0
    estimated_size: int = 0
    _last_size_ts: float = 0.0
    _preview_scanned: bool = False
    stop_event: threading.Event = field(default_factory=threading.Event)
    thread: Optional[threading.Thread] = None
    # ★ 修复：日志读写加锁（对齐双语版，worker 与刷新线程并发访问）
    _lock: threading.RLock = field(
        default_factory=threading.RLock, repr=False, compare=False
    )

    def to_dict(self):
        with self._lock:
            log_snap = list(self.log[-60:])
            files_snap = list(self.output_files or [])
            return {
                "task_id": self.task_id,
                "kind": self.kind,
                "src_name": self.src_name,
                "out_dir": self.out_dir,
                "created_at": self.created_at,
                "status": self.status if self.status != "stopping" else "paused",
                "current": self.current,
                "total": self.total,
                "label": self.label,
                "log": log_snap,
                "output_files": files_snap,
                "error": self.error,
                "current_size": self.current_size,
                "estimated_size": self.estimated_size,
            }

    def log_msg(self, m):
        # ★ 优化：加时间戳
        with self._lock:
            ts = time.strftime(hw_core._LOG_TS_FMT)
            self.log.append(f"[{ts}] {m}")
            if len(self.log) > 200:
                del self.log[:-200]

    def log_tail(self, n=40):
        """★ 修复：锁内取日志尾部快照，供刷新线程安全读取。"""
        with self._lock:
            return "\n".join(self.log[-n:])


class TaskManager:
    def __init__(self):
        self._tasks = {}
        self._lock = threading.Lock()

    def create(self, kind, src_path, src_name, out_dir, work_dir):
        tid = uuid.uuid4().hex[:8]
        t = TaskState(task_id=tid, kind=kind, src_path=src_path,
                      src_name=src_name, out_dir=out_dir, work_dir=work_dir)
        with self._lock:
            self._tasks[tid] = t
        return t

    def get(self, tid):
        with self._lock:
            return self._tasks.get(tid)

    def all_sorted(self):
        with self._lock:
            return sorted(self._tasks.values(), key=lambda t: -t.created_at)

    def running(self):
        with self._lock:
            return [t for t in self._tasks.values()
                    if t.status in ("queued", "running", "stopping")]

    def find_active_by_out_dir(self, out_dir):
        target = os.path.normcase(os.path.abspath(out_dir))
        with self._lock:
            for t in self._tasks.values():
                if os.path.normcase(os.path.abspath(t.out_dir)) == target and \
                   t.status in ("queued", "running", "stopping"):
                    return t
        return None

    def load_from_disk(self, d):
        tid = d.get("task_id")
        if not tid:
            return
        with self._lock:
            if tid in self._tasks:
                return
        t = TaskState(
            task_id=tid,
            kind=d.get("kind", "pdf"),
            src_path="",
            src_name=d.get("src_name", ""),
            out_dir=d.get("out_dir", ""),
            work_dir=os.path.join(d.get("out_dir", ""), "_work"),
            created_at=d.get("created_at", time.time()),
            status="paused",
            current=d.get("current", 0),
            total=d.get("total", 0),
            label=d.get("label", ""),
            log=d.get("log", []),
            output_files=d.get("output_files", []) or [],
            error=d.get("error", ""),
            current_size=d.get("current_size", 0),
            estimated_size=d.get("estimated_size", 0),
        )
        with self._lock:
            self._tasks[tid] = t


MANAGER = TaskManager()


def save_state(task, force=False):
    if task is None:
        return
    now = time.monotonic()
    with _STATE_SAVE_LOCK:
        last = _STATE_SAVE_TS.get(task.task_id)
        if not force and last is not None and \
           (now - last) < hw_core.STATE_SAVE_INTERVAL:
            return
        _STATE_SAVE_TS[task.task_id] = now

    try:
        path = os.path.join(task.work_dir, "state.json")
        os.makedirs(task.work_dir, exist_ok=True)
        with hw_core._IO_LOCK:
            tmp = path + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(task.to_dict(), f, ensure_ascii=False, indent=2)
            os.replace(tmp, path)
    except Exception as e:
        # ★ 修复：静默失败 → 至少可见
        print(f"[warn] save_state 保存失败：{e}")


def scan_all_states():
    root = os.path.abspath(hw_core.RESULT_ROOT)
    if not os.path.isdir(root):
        return
    for name in os.listdir(root):
        p = os.path.join(root, name, "_work", "state.json")
        if os.path.isfile(p):
            try:
                with open(p, "r", encoding="utf-8") as f:
                    MANAGER.load_from_disk(json.load(f))
            except Exception:
                pass


def cleanup_orphan_files():
    """★ 修复：清理崩溃遗留的临时/备用 PDF
    （移植自双语版，额外兼容旧命名 _new_<时间戳>.pdf）。"""
    root = os.path.abspath(hw_core.RESULT_ROOT)
    if not os.path.isdir(root):
        return 0
    removed = 0
    for name in os.listdir(root):
        d = os.path.join(root, name)
        if not os.path.isdir(d):
            continue
        for f in os.listdir(d):
            fp = os.path.join(d, f)
            try:
                if time.time() - os.path.getmtime(fp) < 86400:
                    continue
            except Exception:
                continue
            if (f.endswith(".tmp.pdf") or f.endswith(".bak.pdf")
                    or ("_new_" in f and f.endswith(".pdf"))):
                try:
                    os.remove(fp)
                    removed += 1
                    print(f"   [clean] {name}/{f}")
                except Exception:
                    pass
    return removed


def _task_disk_size(task):
    total = 0
    for f in list(task.output_files or []):
        try:
            if f and os.path.exists(f):
                total += os.path.getsize(f)
        except Exception:
            pass
    return total


def refresh_task_size(task, force=False):
    if task is None:
        return
    now = time.time()
    if not force and (now - task._last_size_ts) < hw_core.SIZE_REFRESH_INTERVAL:
        return
    task._last_size_ts = now
    try:
        task.current_size = _task_disk_size(task)
    except Exception as e:
        # ★ 修复：静默失败 → 至少可见
        print(f"[warn] refresh_task_size 失败：{e}")


# ============================================================
# PDF worker
# ============================================================
def pdf_worker(task, paths, real_key, model, trial,
               subject="general", extra_prompt="",
               ans_position="inside", parse_only=False,
               trial_pages=5):
    # ★ 修复：base_url 进 .env；设 timeout + max_retries=0（避免 SDK
    # 默认 600s 超时 × 内置重试 × 自写重试叠加，停止信号无法打断）
    client = OpenAI(api_key=real_key, base_url=hw_core.DEEPSEEK_BASE_URL,
                    timeout=hw_core.API_TIMEOUT, max_retries=0)
    cache = hw_core.load_json_file(paths["cache_file"], {})
    done_pages = hw_core.load_progress_file(paths["progress_file"])

    try:
        src_sz = os.path.getsize(task.src_path)
    except Exception:
        src_sz = 300 * 1024
    task.estimated_size = int(src_sz * 1.3)
    if task.estimated_size:
        task.log_msg(f"📦 预估总产出 ≈ {hw_core.fmt_size(task.estimated_size)}")

    if parse_only:
        task.log_msg("📄 仅解析模式：不会修改原 PDF，仅输出 analysis.md")
    if subject != "general":
        task.log_msg(f"📚 学科微调：{subject}")

    pdf_exists = os.path.exists(paths["solved_pdf"])
    # ★ 优化：仅解析模式下不依赖 solved_pdf 存在
    if not parse_only and done_pages and pdf_exists:
        try:
            with open(paths["solved_pdf"], "rb") as f:
                data = f.read()
            doc = fitz.open(stream=data, filetype="pdf")
            task.log_msg(f"📂 从已解题 PDF 续传（进度记录：{len(done_pages)} 页）")
        except Exception as e:
            task.log_msg(f"⚠️ 打开旧结果失败，从头开始：{e}")
            doc = fitz.open(task.src_path)
            done_pages = set()
            hw_core.save_progress_file(paths["progress_file"], done_pages)
    else:
        doc = fitz.open(task.src_path)
        if done_pages and not pdf_exists and not parse_only:
            task.log_msg("⚠️ 有进度记录但缺结果 PDF，从头开始")
            done_pages = set()
            hw_core.save_progress_file(paths["progress_file"], done_pages)

    total = len(doc)
    # ★ 修复：试解页数可选（3/5/10/20），不再固定 5 页
    limit = min(trial_pages, total) if trial else total

    task.status = "running"
    task.total = limit
    task.label = f"PDF · 目标前 {limit} 页"
    task.current = len([p for p in done_pages if p <= limit])
    task.log_msg(f"✅ PDF 共 {total} 页，本次目标 {limit} 页，已完成 {task.current} 页")
    save_state(task, force=True)

    if not parse_only and pdf_exists:
        try:
            task.preview_images = render_preview_only(
                paths["solved_pdf"], paths, task)
            task.output_files = [paths["solved_pdf"]]
            refresh_task_size(task, force=True)
            save_state(task, force=True)
        except Exception as e:
            task.log_msg(f"⚠️ 初始预览失败：{e}")

    error_msg = None
    markdown_items = []   # (page_no, [(src, ans, sol), ...])

    for pno in range(total):
        if task.stop_event.is_set():
            task.log_msg("⏸ 停止信号，结束当前循环")
            break
        page_num = pno + 1
        if page_num > limit:
            break
        if page_num in done_pages and not parse_only:
            continue

        page = doc[pno]
        blocks = [b for b in page.get_text("blocks")
                  if b[6] == 0 and b[4].strip()]

        if not blocks:
            done_pages.add(page_num)
            hw_core.save_progress_file(paths["progress_file"], done_pages)
            task.current = len([p for p in done_pages if p <= limit])
            task.label = f"已跳过空白/图片页 {page_num}"
            save_state(task)
            continue

        try:
            results = solve_page(client, model, blocks, cache,
                                 paths["cache_file"],
                                 stop_event=task.stop_event,
                                 subject=subject,
                                 extra_prompt=extra_prompt)

            # 收集到 markdown
            page_qa = []
            for i, b in enumerate(blocks):
                res = results.get(i)
                if not res:
                    continue
                src = (b[4] or "").strip()
                a = (res.get("ans") or "").strip()
                s = (res.get("sol") or "").strip()
                if a or s:
                    page_qa.append((src, a, s))
            if page_qa:
                markdown_items.append((page_num, page_qa))

            # 只有非仅解析模式才写回 PDF
            if not parse_only:
                apply_solution_to_page(page, blocks, results,
                                       ans_position=ans_position,
                                       warn=lambda m: task.log_msg(m))
        except RuntimeError as e:
            error_msg = str(e)
            break
        except Exception as e:
            error_msg = f"未知错误：{e}"
            break

        done_pages.add(page_num)
        hw_core.save_progress_file(paths["progress_file"], done_pages)
        task.current = len([p for p in done_pages if p <= limit])
        task.label = f"已解 {task.current}/{limit} 页"
        task.log_msg(f"✅ 第 {page_num} 页完成")
        save_state(task, force=(page_num % 5 == 0))

        # 仅解析模式不用每次落盘 PDF
        if parse_only:
            continue

        should_save = (page_num <= hw_core.PREVIEW_PAGES) or \
                      (page_num % hw_core.CHECKPOINT_EVERY == 0)
        if should_save:
            try:
                saved = hw_core.safe_save_pdf(doc, paths["solved_pdf"])
                if saved != paths["solved_pdf"]:
                    task.log_msg(
                        f"⚠️ 原文件被占用，已保存到备用路径："
                        f"{os.path.basename(saved)}"
                    )
                    paths["solved_pdf"] = saved
                task.log_msg(f"💾 已落盘（前 {page_num} 页）")
                if page_num <= hw_core.PREVIEW_PAGES:
                    task.preview_images = render_preview_only(
                        paths["solved_pdf"], paths, task)
                task.output_files = [paths["solved_pdf"]]
                refresh_task_size(task, force=True)
                save_state(task, force=True)
            except Exception as e:
                task.log_msg(f"⚠️ 落盘失败：{e}")

    save_failed = False
    if not parse_only:
        try:
            saved = hw_core.safe_save_pdf(doc, paths["solved_pdf"])
            if saved != paths["solved_pdf"]:
                task.log_msg(
                    f"⚠️ 原文件被占用，已保存到备用路径："
                    f"{os.path.basename(saved)}"
                )
                paths["solved_pdf"] = saved
        except Exception as e:
            save_failed = True
            error_msg = error_msg or f"保存结果 PDF 失败：{e}"
    try:
        doc.close()
    except Exception:
        pass

    if not parse_only:
        if save_failed or not os.path.exists(paths["solved_pdf"]):
            task.status = "error"
            task.error = error_msg or "结果 PDF 保存失败"
            task.label = "出错（保存失败）"
            task.log_msg(f"❌ {task.error}")
            refresh_task_size(task, force=True)
            save_state(task, force=True)
            return

    if not done_pages:
        task.status = "error"
        task.error = "本次没有任何页面成功处理"
        task.log_msg("⚠️ " + task.error)
        refresh_task_size(task, force=True)
        save_state(task, force=True)
        return

    # 仅解析模式：写 markdown
    if parse_only:
        try:
            book_title = os.path.splitext(task.src_name)[0]
            md = build_solution_markdown(book_title, markdown_items)
            with open(paths["analysis_md"], "w", encoding="utf-8") as f:
                f.write(md)
            task.output_files = [paths["analysis_md"]]
            task.log_msg(f"📝 已输出解析报告：{paths['analysis_md']}")
        except Exception as e:
            task.log_msg(f"⚠️ 写 markdown 失败：{e}")
    else:
        try:
            task.preview_images = render_preview_only(
                paths["solved_pdf"], paths, task)
        except Exception:
            pass
        task.output_files = [paths["solved_pdf"]]

    target_pages = set(range(1, limit + 1))
    completed = (target_pages & done_pages == target_pages) and not error_msg

    if completed:
        task.status = "done"
        task.label = "全部完成"
        task.log_msg(f"🎉 全部完成！共 {len(done_pages)} 页")
    elif task.stop_event.is_set():
        task.status = "paused"
        task.label = f"已暂停（半成品），共 {len(done_pages)} 页"
        task.log_msg("🛑 已暂停，下次从未完成的页继续")
    elif error_msg:
        task.status = "error"
        task.error = error_msg
        task.label = "出错（半成品）"
        task.log_msg(f"❌ {error_msg}")
    else:
        task.status = "paused"
        task.label = "半成品"
    refresh_task_size(task, force=True)
    save_state(task, force=True)


# ============================================================
# Word / PPT worker
# ============================================================
def _docx_insert_after(para, text):
    new_p = OxmlElement("w:p")
    para._p.addnext(new_p)
    new_para = Paragraph(new_p, para._parent)
    if text:
        run = new_para.add_run(text)
        if para.runs:
            try:
                src_rpr = para.runs[0]._element.find(qn('w:rPr'))
                if src_rpr is not None:
                    run._element.insert(0, copy.deepcopy(src_rpr))
            except Exception:
                pass
    try:
        new_para.style = para.style
    except Exception:
        pass
    try:
        pPr = new_para._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "EEF2FF")
        pPr.append(shd)
    except Exception:
        pass
    return new_para


def _pptx_insert_after(para, text):
    """★ 修复：PPT 答案/解析插成独立新段落（对齐 docx 的 _docx_insert_after），
    不再用 \\v 拼进原段落（部分播放器渲染异常）。"""
    new_p = copy.deepcopy(para._p)
    for r in new_p.findall(qn('a:r')):
        new_p.remove(r)
    para._p.addnext(new_p)
    r = new_p.makeelement(qn('a:r'), {})
    # 保持 <a:p> 子元素顺序合法：run 必须在 pPr 之后
    pPr = new_p.find(qn('a:pPr'))
    if pPr is not None:
        pPr.addnext(r)
    else:
        new_p.insert(0, r)
    if para.runs:
        try:
            src_rpr = para.runs[0]._element.find(qn('a:rPr'))
            if src_rpr is not None:
                r.append(copy.deepcopy(src_rpr))
        except Exception:
            pass
    t_el = r.makeelement(qn('a:t'), {})
    r.append(t_el)
    t_el.text = text
    return new_p


def _iter_pptx_shapes(shapes):
    for shape in shapes:
        if MSO_SHAPE_TYPE is not None and \
                getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            try:
                yield from _iter_pptx_shapes(shape.shapes)
                continue
            except Exception:
                pass
        yield shape


def docx_worker(task, paths, real_key, model,
                subject="general", extra_prompt=""):
    # ★ 修复：base_url 进 .env；设 timeout + max_retries=0（避免 SDK
    # 默认 600s 超时 × 内置重试 × 自写重试叠加，停止信号无法打断）
    client = OpenAI(api_key=real_key, base_url=hw_core.DEEPSEEK_BASE_URL,
                    timeout=hw_core.API_TIMEOUT, max_retries=0)
    cache = hw_core.load_json_file(paths["cache_file"], {})

    try:
        src_sz = os.path.getsize(task.src_path)
    except Exception:
        src_sz = 300 * 1024
    task.estimated_size = int(src_sz * 1.5)
    if task.estimated_size:
        task.log_msg(f"📦 预估总产出 ≈ {hw_core.fmt_size(task.estimated_size)}")

    try:
        wdoc = Document(task.src_path)
    except Exception as e:
        task.status = "error"
        task.error = f"打开 Word 失败：{e}"
        task.log_msg(f"❌ {task.error}")
        save_state(task, force=True)
        return

    targets = [p for p in wdoc.paragraphs if p.text.strip()]

    _seen_cells = set()
    for table in wdoc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in _seen_cells:
                    continue
                _seen_cells.add(cell._tc)
                for p in cell.paragraphs:
                    if p.text.strip():
                        targets.append(p)

    total = len(targets)
    task.status = "running"
    task.total = total
    task.current = 0
    task.label = f"Word · 共 {total} 段"
    task.log_msg(f"📘 Word 已打开，共 {total} 段（含表格）")
    save_state(task, force=True)

    error_msg = None
    done = 0
    solved_count = 0
    preview_pairs = []

    # ★ 修复：批量解题（20 段/次 API 调用），替代逐段调用
    all_texts = [p.text.strip() for p in targets]
    try:
        results = solve_batch_office(
            client, model, all_texts, cache, paths["cache_file"],
            prefix="w", stop_event=task.stop_event, task=task,
            subject=subject, extra_prompt=extra_prompt,
        )
    except RuntimeError as e:
        error_msg = str(e)
        results = None

    if results is not None:
        for para, text, res in zip(targets, all_texts, results):
            if res is None:
                break  # 停止信号，本批未处理完
            ans = (res.get("ans") or "").strip()
            sol = (res.get("sol") or "").strip()
            if ans or sol:
                ins = []
                if ans:
                    ins.append(f"【答案】{ans}")
                if sol:
                    ins.append(f"【解析】{sol}")
                rendered = "\n".join(ins)
                _docx_insert_after(para, rendered)
                solved_count += 1
                if len(preview_pairs) < hw_core.PREVIEW_PARAS:
                    preview_pairs.append((text, rendered))
                    task.preview_html = _build_docx_preview_html(preview_pairs)

            done += 1
            task.current = done
            task.label = f"Word {done}/{total}（已解 {solved_count}）"
            if done % 5 == 0:
                save_state(task)

    try:
        wdoc.save(paths["office_out"])
        task.output_files = [paths["office_out"]]
        task.log_msg(f"✅ 已保存：{paths['office_out']}")
    except Exception as e:
        error_msg = error_msg or f"保存 Word 失败：{e}"

    if task.stop_event.is_set():
        task.status = "paused"
        task.label = f"已暂停 {done}/{total}"
    elif error_msg:
        task.status = "error"
        task.error = error_msg
    else:
        task.status = "done"
        task.label = "全部完成"
        task.log_msg(
            f"🎉 Word 解题完成！共处理 {done} 段，其中 {solved_count} 段有答案")
    refresh_task_size(task, force=True)
    save_state(task, force=True)


def pptx_worker(task, paths, real_key, model,
                subject="general", extra_prompt=""):
    # ★ 修复：base_url 进 .env；设 timeout + max_retries=0（避免 SDK
    # 默认 600s 超时 × 内置重试 × 自写重试叠加，停止信号无法打断）
    client = OpenAI(api_key=real_key, base_url=hw_core.DEEPSEEK_BASE_URL,
                    timeout=hw_core.API_TIMEOUT, max_retries=0)
    cache = hw_core.load_json_file(paths["cache_file"], {})

    try:
        src_sz = os.path.getsize(task.src_path)
    except Exception:
        src_sz = 300 * 1024
    task.estimated_size = int(src_sz * 1.5)
    if task.estimated_size:
        task.log_msg(f"📦 预估总产出 ≈ {hw_core.fmt_size(task.estimated_size)}")

    try:
        prs = Presentation(task.src_path)
    except Exception as e:
        task.status = "error"
        task.error = f"打开 PPT 失败：{e}"
        task.log_msg(f"❌ {task.error}")
        save_state(task, force=True)
        return

    targets = []
    for si, slide in enumerate(prs.slides):
        for shape in _iter_pptx_shapes(slide.shapes):
            if not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                txt = "".join(r.text for r in para.runs)
                if txt.strip():
                    targets.append((si, para))

    total = len(targets)
    task.status = "running"
    task.total = total
    task.current = 0
    task.label = f"PPT · 共 {total} 段"
    task.log_msg(f"📊 PPT 已打开，共 {total} 段（含组合形状）")
    save_state(task, force=True)

    error_msg = None
    done = 0
    solved_count = 0
    preview_pairs = []

    # ★ 修复：批量解题（20 段/次 API 调用），替代逐段调用
    all_texts = ["".join(r.text for r in para.runs).strip()
                 for _si, para in targets]
    try:
        results = solve_batch_office(
            client, model, all_texts, cache, paths["cache_file"],
            prefix="p", stop_event=task.stop_event, task=task,
            subject=subject, extra_prompt=extra_prompt,
        )
    except RuntimeError as e:
        error_msg = str(e)
        results = None

    if results is not None:
        for (si, para), text, res in zip(targets, all_texts, results):
            if res is None:
                break  # 停止信号，本批未处理完
            ans = (res.get("ans") or "").strip()
            sol = (res.get("sol") or "").strip()
            if ans or sol:
                parts = []
                if ans:
                    parts.append(f"【答案】{ans}")
                if sol:
                    parts.append(f"【解析】{sol}")
                rendered = "\n".join(parts)
                _pptx_insert_after(para, rendered)
                solved_count += 1
                if len(preview_pairs) < hw_core.PREVIEW_PARAS:
                    preview_pairs.append((si + 1, text, rendered))
                    task.preview_html = _build_pptx_preview_html(preview_pairs)

            done += 1
            task.current = done
            task.label = f"PPT {done}/{total}（第 {si+1} 张，已解 {solved_count}）"
            if done % 5 == 0:
                save_state(task)

    try:
        prs.save(paths["office_out"])
        task.output_files = [paths["office_out"]]
        task.log_msg(f"✅ 已保存：{paths['office_out']}")
    except Exception as e:
        error_msg = error_msg or f"保存 PPT 失败：{e}"

    if task.stop_event.is_set():
        task.status = "paused"
        task.label = f"已暂停 {done}/{total}"
    elif error_msg:
        task.status = "error"
        task.error = error_msg
    else:
        task.status = "done"
        task.label = "全部完成"
        task.log_msg(
            f"🎉 PPT 解题完成！共处理 {done} 段，其中 {solved_count} 段有答案")
    refresh_task_size(task, force=True)
    save_state(task, force=True)
