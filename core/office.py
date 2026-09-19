# -*- coding: utf-8 -*-
"""Office 工具：docx/pptx 段落替换、批量翻译。"""

from core import config
from core.utils import h, save_json_file
from core.api import (build_simple_system, _api_call, call_simple_api,
                      cache_prefix, _is_fatal_api_error)
from core.pdf_pipeline import parse_marked

# Word / PPT 依赖为可选；未安装时这些名字为 None，
# 相关 worker 仅在 HAS_OFFICE 时才会被调用。
if config.HAS_OFFICE:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    from pptx import Presentation
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except Exception:
        MSO_SHAPE_TYPE = None
else:
    Document = qn = OxmlElement = Paragraph = Presentation = None
    MSO_SHAPE_TYPE = None


# ============================================================
# Office
# ============================================================

def _docx_replace_para_text(para, new_text):
    if para.runs:
        para.runs[0].text = new_text
        for r in para.runs[1:]:
            r.text = ""
    else:
        para.add_run(new_text)


def _docx_insert_after(para, text):
    new_p = OxmlElement("w:p")
    para._p.addnext(new_p)
    new_para = Paragraph(new_p, para._parent)
    if text:
        new_para.add_run(text)
    try:
        new_para.style = para.style
    except Exception:
        pass
    try:
        pPr = new_para._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F5F1E6")
        pPr.append(shd)
    except Exception:
        pass
    return new_para


def _pptx_set_para_text(para, new_text):
    if para.runs:
        para.runs[0].text = new_text
        for r in para.runs[1:]:
            r.text = ""
    else:
        try:
            run = para.add_run()
            run.text = new_text
        except Exception:
            try:
                para.text = new_text
            except Exception:
                pass


def _iter_pptx_shapes(shapes):
    for shape in shapes:
        if MSO_SHAPE_TYPE is not None and \
                getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            try:
                yield from _iter_pptx_shapes(shape.shapes)
                continue
            except Exception:
                pass
        yield shape


def translate_batch_office(client, model, texts, cache, cache_file,
                           target_lang="zh-CN", stop_event=None, task=None,
                           domain="general", extra_prompt=""):
    pref = cache_prefix(target_lang)
    results = [None] * len(texts)
    failed = set()
    any_success = False

    def ckey(t):
        return f"t_{pref}" + h(t)

    pending = []
    for i, t in enumerate(texts):
        k = ckey(t)
        if k in cache and isinstance(cache[k], str) and cache[k].strip():
            results[i] = cache[k]
            any_success = True
        else:
            pending.append((i, t))

    for start in range(0, len(pending), config.OFFICE_BATCH_SIZE):
        if stop_event is not None and stop_event.is_set():
            break
        chunk = pending[start:start + config.OFFICE_BATCH_SIZE]
        marked = "\n\n".join(f"[[B{j}]] {t}" for j, (_, t) in enumerate(chunk))
        try:
            system = build_simple_system(target_lang, domain, extra_prompt)
            raw = _api_call(client, model, system, marked, stop_event=stop_event)
            parsed = parse_marked(raw, len(chunk))
            for j, (gi, t) in enumerate(chunk):
                tr = (parsed.get(j) or "").strip()
                if tr:
                    results[gi] = tr
                    cache[ckey(t)] = tr
                    any_success = True
                else:
                    failed.add(gi)
        except RuntimeError as e:
            if _is_fatal_api_error(str(e)):
                raise
            if task:
                task.log_msg(f"⚠️ 批量翻译失败（{e}），改为逐条重试")
            for gi, t in chunk:
                if stop_event is not None and stop_event.is_set():
                    break
                try:
                    tr = call_simple_api(client, model, t, target_lang,
                                         stop_event=stop_event,
                                         domain=domain,
                                         extra_prompt=extra_prompt)
                    results[gi] = tr
                    cache[ckey(t)] = tr
                    any_success = True
                except RuntimeError as e2:
                    if _is_fatal_api_error(str(e2)):
                        raise
                    failed.add(gi)
                    if task:
                        task.log_msg(f"⚠️ 段落 #{gi} 翻译失败：{e2}")
                    continue

        save_json_file(cache_file, cache)

    if pending and not any_success and not (
        stop_event is not None and stop_event.is_set()
    ):
        raise RuntimeError("所有段落翻译均失败，请检查 API Key / 余额 / 网络")

    # ★ 修复：失败段落保持 None（不再用原文冒充译文），由调用方跳过并标记
    fail_count = sum(1 for r in results if r is None)

    return results, fail_count
